"""AMFI's total expense ratios, every fund, every day. DECISIONS V1-78.

AMFI's `ter-of-mf-schemes` page is a Next.js app like its fund-size page (S3):
its data comes from `/api/populate-te-rdata-revised`, neither authenticated nor
challenged, and robots.txt disallows only `/admin/`, `/login/` and `/search/`.
As JSON it pages at 100 rows -- ~500 requests for one month -- so this reads
the page's own "download" instead: `excel=true`, the whole month in one
workbook (~4 MB, ~20 s for AMFI to build).

**One row per fund per day**, the Regular and Direct plans side by side, each
with its base expense ratio, the three trading-cost components and the total.
Fund-level names ("HDFC Flexi Cap Fund"), no ISIN: `jobs/fetch_ter.py` joins
them to share classes through `family_key`, as holdings are (V1-37).

**Percent, as published** (0.62 is 0.62% a year). Excel cells arrive as
floats; each is read through its shortest repr and held to four places, the
precision AMFI's own JSON carries ("0.6200").
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

URL = (
    "https://www.amfiindia.com/api/populate-te-rdata-revised"
    "?MF_ID=All&Month={month}&strCat=-1&strType=-1&excel=true"
)
SHEET = "TER_Revised"
#: The columns this reader depends on, by the header AMFI writes.
NAME, CATEGORY, DAY = "Scheme Name", "Scheme Category", "TER Date"
PLANS = {
    "regular": ("Regular Plan - Total TER (%)",
                "Regular Plan - Base Expense Ratio (BER) (%)"),
    "direct": ("Direct Plan - Total TER (%)",
               "Direct Plan - Base Expense Ratio (BER) (%)"),
}
PLACES = Decimal("0.0001")


class TerPayloadError(ValueError):
    """The workbook is not the shape this reader was written against."""


class NothingPublished(TerPayloadError):
    """A well-formed workbook with no rows: a month AMFI has not started."""


@dataclass(frozen=True)
class FundTer:
    """One fund's one plan on one day, as published."""

    fund_name: str
    category: str
    day: date
    plan: str
    total: Decimal
    base: Decimal | None


def month_url(on: date) -> str:
    return URL.format(month=f"{on:%m-%Y}")


def _percent(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    return Decimal(str(value)).quantize(PLACES)


def parse_ter(payload: bytes) -> list[FundTer]:
    """Every fund, plan and day the workbook lists. A plan with no total (a
    fund with only one plan) is absent, never zero."""
    try:
        book = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:  # zipfile, XML: anything that is not a workbook
        raise TerPayloadError(f"not a workbook: {exc}") from exc
    if SHEET not in book.sheetnames:
        raise TerPayloadError(f"no {SHEET!r} sheet; found {book.sheetnames}")
    rows = book[SHEET].iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(rows, ())]
    wanted = [NAME, CATEGORY, DAY, *(c for pair in PLANS.values() for c in pair)]
    if missing := [c for c in wanted if c not in header]:
        raise TerPayloadError(f"columns missing: {missing}")
    at = {c: header.index(c) for c in wanted}

    out: list[FundTer] = []
    for row in rows:
        # The sheet ends with the page's disclaimer, one sentence per row.
        if len(row) < len(header) or not isinstance(row[at[DAY]], datetime):
            continue
        name = str(row[at[NAME]] or "").strip()
        if not name:
            continue
        for plan, (total_col, base_col) in PLANS.items():
            total = _percent(row[at[total_col]])
            if total is None:
                continue
            out.append(FundTer(
                fund_name=name,
                category=str(row[at[CATEGORY]] or "").strip(),
                day=row[at[DAY]].date(),
                plan=plan,
                total=total,
                base=_percent(row[at[base_col]]),
            ))
    if not out:
        raise NothingPublished("the workbook lists no expense ratios")
    return out


__all__ = ["URL", "FundTer", "NothingPublished", "TerPayloadError", "month_url",
           "parse_ter"]
