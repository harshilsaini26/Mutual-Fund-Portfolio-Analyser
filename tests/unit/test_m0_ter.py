"""Expense ratios: AMFI's workbook to `scheme_ter` to a fund's TER. DECISIONS V1-78.

The fixture is cut from AMFI's real August 2026 workbook: its header, three
funds on the month's last two days, and the disclaimer the sheet ends with.
"""

from __future__ import annotations

import sqlite3
from collections.abc import Callable, Iterator
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from jobs import fetch_ter
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from src.common.decimals import connect
from src.common.types import SchemeId
from src.m0_data.fetch.amfi_ter import (
    FundTer,
    NothingPublished,
    TerPayloadError,
    month_url,
    parse_ter,
)
from src.m0_data.providers.warehouse import WarehouseMarketDataProvider
from src.m2_fund.peers import peer_context

from tests.conftest import migrated

FIXTURE = Path(__file__).parents[1] / "fixtures" / "m0" / "amfi_ter_2026-08_sample.xlsx"
AUG_31 = date(2026, 8, 31)


def _book(edit: Callable[[Worksheet], None] | None = None) -> bytes:
    """The fixture, optionally changed by `edit(sheet)`, as bytes."""
    book = load_workbook(FIXTURE)
    if edit is not None:
        edit(book["TER_Revised"])
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def test_each_fund_plan_and_day_is_read_as_published() -> None:
    rows = parse_ter(FIXTURE.read_bytes())
    hdfc = [r for r in rows if r.fund_name == "HDFC Flexi Cap Fund" and r.day == AUG_31]
    assert {r.plan: r.total for r in hdfc} == {
        "regular": Decimal("1.3700"), "direct": Decimal("0.7700"),
    }
    assert all(r.base is not None and r.base < r.total for r in hdfc)
    # Three funds, two days, two plans; the disclaimer rows are not funds.
    assert len(rows) == 12


def test_a_plan_with_no_figure_is_absent_not_zero() -> None:
    def drop_regular(sheet: Worksheet) -> None:
        sheet.cell(row=2, column=10).value = None

    rows = parse_ter(_book(drop_regular))
    assert len(rows) == 11
    assert all(r.total > 0 for r in rows)


def test_a_workbook_of_another_shape_is_refused() -> None:
    with pytest.raises(TerPayloadError, match="not a workbook"):
        parse_ter(b"<html>maintenance</html>")

    def rename(sheet: Worksheet) -> None:
        sheet.cell(row=1, column=15).value = "Direct TER"

    with pytest.raises(TerPayloadError, match="columns missing"):
        parse_ter(_book(rename))

    def empty(sheet: Worksheet) -> None:
        sheet.delete_rows(2, 100)

    with pytest.raises(NothingPublished):
        parse_ter(_book(empty))


def test_a_month_is_named_the_way_amfi_names_it() -> None:
    assert "Month=08-2026" in month_url(AUG_31)
    assert fetch_ter.last_month(date(2026, 9, 25)) == AUG_31
    assert fetch_ter.last_month(date(2026, 1, 1)) == date(2025, 12, 31)


def _ter(name: str, day: int, total: str, plan: str = "direct") -> FundTer:
    return FundTer(name, "Equity Scheme - ELSS", date(2026, 8, day), plan,
                   Decimal(total), None)


def test_the_last_day_wins_and_two_spellings_are_one_fund() -> None:
    figures, clashes = fetch_ter.month_end([
        _ter("Axis ELSS Tax Saver Fund", 30, "0.70"),
        _ter("Axis ELSS- Tax Saver Fund", 31, "0.72"),
        _ter("Axis ELSS Tax Saver Fund", 29, "0.99"),
    ])
    assert clashes == 0
    assert figures[("axis elss tax saver fund", "direct")].total == Decimal("0.72")


def test_two_different_figures_on_the_last_day_are_dropped_not_picked() -> None:
    figures, clashes = fetch_ter.month_end([
        _ter("Axis ELSS Tax Saver Fund", 31, "0.72"),
        _ter("Axis ELSS- Tax Saver Fund", 31, "0.80"),
        _ter("Axis ELSS Tax Saver Fund", 31, "1.50", plan="regular"),
    ])
    assert clashes == 1
    assert list(figures) == [("axis elss tax saver fund", "regular")]


@pytest.fixture
def warehouse(tmp_path: Path) -> Iterator[sqlite3.Connection]:
    db = tmp_path / "w.db"
    migrated(db)
    conn = connect(str(db))
    conn.execute("INSERT INTO raw_file (file_id, source_id, fetched_at, storage_path,"
                  " byte_size) VALUES ('f1', 'S14', '2026-09-25', '/x', 0)")
    for sid, family, plan, option, amc in (
        ("INF179K01UT0", "hdfc flexi cap fund", "direct", "growth", "hdfc"),
        ("INF179K01VL5", "hdfc flexi cap fund", "direct", "idcw_payout", "hdfc"),
        ("INF179K01608", "hdfc flexi cap fund", "regular", "growth", "hdfc"),
        # One name, two fund houses: no figure can say which it belongs to.
        ("INF000A00001", "twin fund", "direct", "growth", "amc_a"),
        ("INF000B00001", "twin fund", "direct", "growth", "amc_b"),
    ):
        conn.execute(
            "INSERT INTO scheme (scheme_id, scheme_name, plan, option, amc_id,"
            " scheme_family, status, last_seen) VALUES (?,?,?,?,?,?, 'active', ?)",
            (sid, family.title(), plan, option, amc, family, date(2026, 9, 24)),
        )
    conn.commit()
    yield conn
    conn.close()


def test_a_fund_figure_reaches_every_share_class_of_its_plan(
    warehouse: sqlite3.Connection,
) -> None:
    conn = warehouse
    figures, _ = fetch_ter.month_end(parse_ter(FIXTURE.read_bytes()))
    figures[("twin fund", "direct")] = _ter("Twin Fund", 31, "0.50")
    joined, classes, added = fetch_ter.load(conn, figures, "f1")
    # HDFC's two plans join (three share classes); 360 ONE and Axis are not in
    # this warehouse, and the twin fund's family spans two houses.
    assert (joined, classes, added) == (2, 3, 3)

    market = WarehouseMarketDataProvider(conn)
    assert market.ter(SchemeId("INF179K01VL5"), AUG_31) == Decimal("0.7700")
    assert market.ter(SchemeId("INF179K01608"), AUG_31) == Decimal("1.3700")
    with pytest.raises(KeyError):
        market.ter(SchemeId("INF000A00001"), AUG_31)
    with pytest.raises(KeyError):  # nothing was in force before the 31st
        market.ter(SchemeId("INF179K01UT0"), date(2026, 8, 30))


def test_loading_again_adds_nothing_and_a_correction_is_a_revision(
    warehouse: sqlite3.Connection,
) -> None:
    conn = warehouse
    figures, _ = fetch_ter.month_end(parse_ter(FIXTURE.read_bytes()))
    fetch_ter.load(conn, figures, "f1")
    assert fetch_ter.load(conn, figures, "f1")[2] == 0

    key = ("hdfc flexi cap fund", "direct")
    old = figures[key]
    figures = {key: FundTer(old.fund_name, old.category, old.day, old.plan,
                            Decimal("0.7900"), old.base)}
    assert fetch_ter.load(conn, figures, "f1")[2] == 2
    revisions = conn.execute(
        "SELECT revision FROM scheme_ter WHERE scheme_id = 'INF179K01UT0'"
        " ORDER BY revision").fetchall()
    assert [r[0] for r in revisions] == [1, 2]
    market = WarehouseMarketDataProvider(conn)
    assert market.ter(SchemeId("INF179K01UT0"), AUG_31) == Decimal("0.7900")
    facts = market.scheme_facts(SchemeId("INF179K01UT0"))
    assert facts is not None and facts.ter == Decimal("0.7900")
    assert facts.ter_as_of == AUG_31


def test_the_cheapest_fund_ranks_first_among_its_peers(
    warehouse: sqlite3.Connection,
) -> None:
    conn = warehouse
    for n, ter in enumerate(("0.40", "0.90", "0.60", "1.10", "0.75")):
        sid = f"INF000C0000{n}"
        conn.execute(
            "INSERT INTO scheme (scheme_id, scheme_name, plan, option, amc_id,"
            " sebi_category, status, last_seen) VALUES (?, ?, 'direct', 'growth',"
            " ?, 'Equity Scheme - Flexi Cap Fund', 'active', ?)",
            (sid, f"Fund {n}", f"amc_{n}", date(2026, 9, 24)),
        )
        conn.execute(
            "INSERT INTO scheme_ter (scheme_id, valid_from, total_ter, source_file_id,"
            " ingested_at) VALUES (?, ?, ?, 'f1', ?)",
            (sid, AUG_31, Decimal(ter), datetime(2026, 9, 25)),
        )
    found = peer_context(WarehouseMarketDataProvider(conn), "INF000C00002")
    assert found is not None
    cost = next(r for r in found.ranks if r.metric.key == "ter")
    assert (cost.rank, cost.ranked, cost.value) == (2, 5, Decimal("0.60"))
